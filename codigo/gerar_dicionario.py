"""Extrai dicionario oficial da planilha do INEP (sem gabaritos)."""
from pathlib import Path
import csv
import openpyxl


ARQUIVO_XLSX = Path(
    "dados/censo_escolar/originais/microdados_2025/"
    "dicionário_dados_educação_básica.xlsx"
)
PASTA_CSV = Path('dados/censo_escolar/originais/microdados_2025')
SAIDA = Path('dados/censo_escolar/originais/dicionario_variaveis.txt')

def _limpar(valor):
    """Junta linhas em uma só, para não criar linha em branco dentro do bloco."""
    return ' '.join(str(valor or '').split())


def gerar():
    if not ARQUIVO_XLSX.exists():
        raise FileNotFoundError(f'Coloque a planilha oficial em: {ARQUIVO_XLSX}')
    csvs = sorted(PASTA_CSV.glob('Tabela_Matricula_2025_V2*.csv'))
    if not csvs:
        raise FileNotFoundError(f'CSV de matrículas não encontrado em {PASTA_CSV}')
    with csvs[0].open(encoding='latin-1', newline='') as f:
        colunas = set(next(csv.reader(f, delimiter=';')))
    wb = openpyxl.load_workbook(ARQUIVO_XLSX, read_only=True, data_only=True)
    ws = next((s for s in wb if 'Matrícula' in s.title or 'Matricula' in s.title), None)
    if ws is None:
        raise ValueError('Aba Tabela_de_Matrícula não encontrada')
    # A coluna "Notas Importantes" traz regras do INEP (ex.: o que QT_MAT_BAS não inclui).
    col_notas = 25
    for linha in ws.iter_rows(values_only=True):
        achou = [i for i, c in enumerate(linha) if str(c or '').strip().lower().startswith('notas importantes')]
        if achou:
            col_notas = achou[0]
            break
    linhas = ['DICIONÁRIO OFICIAL INEP - TABELA DE MATRÍCULA',
              f'Origem: {ARQUIVO_XLSX.name}',
              'Descrição e categorias transcritas da planilha original.', '']
    encontrados = set()
    for linha in ws.iter_rows(values_only=True):
        nome = str(linha[1] or '').strip().upper() if len(linha)>1 else ''
        if nome not in colunas or nome in encontrados:
            continue
        encontrados.add(nome)
        descricao = _limpar(linha[2]) if len(linha)>2 else ''
        categoria = '\n'.join(l.strip() for l in str(linha[5] or '').splitlines() if l.strip()) if len(linha)>5 else ''
        notas = _limpar(linha[col_notas]) if len(linha)>col_notas else ''
        bloco = [f'VARIÁVEL: {nome}', f'DESCRIÇÃO: {descricao}', f'CATEGORIAS: {categoria}']
        if notas:
            bloco.append(f'NOTAS: {notas}')
        linhas.extend(bloco + [''])
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text('\n'.join(linhas), encoding='utf-8')
    print(f'{len(encontrados)} variáveis documentadas em {SAIDA}')
    print(f'{len(colunas - encontrados)} colunas do CSV sem descrição correspondente.')

if __name__ == '__main__':
    gerar()
